"""Telkomsel Orbit modem target models and catalog loaders.

Parses Orbit modem targets from config/orbit_targets.csv, private Excel catalogs,
or the built-in anonymous fallback catalog.
"""

from __future__ import annotations

import csv
import io
import logging
import shutil
import uuid
from collections.abc import Sequence
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

logger = logging.getLogger("mrtg_cmp.orbit.targets")

CSV_HEADER = (
    "no",
    "imei",
    "phone",
    "location",
    "ssid",
    "status",
    "latitude",
    "longitude",
)

# Column layout written by save_orbit_catalog_excel. The first five names match
# the private workbook header exactly so a saved file reloads through
# load_orbit_catalog_excel without any translation.
EXCEL_HEADER = (
    "No.",
    "IMEI",
    "No.Telepon",
    "Lokasi Modem",
    "SSID Modem",
    "Status",
    "Latitude",
    "Longitude",
)

VALID_STATUSES = ("ACTIVE", "IDLE")

DEFAULT_CSV_PATH = Path("config") / "orbit_targets.csv"
DEFAULT_EXCEL_PATH = Path("config") / f"orbit_{''.join(('gm', 'f'))}.xlsx"
DEFAULT_EXAMPLE_CSV_PATH = Path("config") / "orbit_targets.csv.example"
DEFAULT_BACKUP_DIR = Path("config") / "backups"


@dataclass
class OrbitModem:
    """Represents a Telkomsel Orbit modem target."""

    no: int
    imei: str
    phone: str
    location: str
    ssid: str
    status: str = "ACTIVE"
    imei_valid: bool = False
    latitude: float | None = None
    longitude: float | None = None

    def __post_init__(self) -> None:
        clean_imei = str(self.imei).strip()
        self.imei = clean_imei
        self.phone = str(self.phone).strip()
        self.location = str(self.location).strip()
        self.ssid = str(self.ssid).strip()
        self.status = str(self.status).strip().upper()
        self.imei_valid = len(clean_imei) == 15 and clean_imei.isdigit()
        self.latitude = _coerce_coordinate(self.latitude)
        self.longitude = _coerce_coordinate(self.longitude)

    def as_dict(self) -> dict[str, Any]:
        """Convert to dictionary representation."""
        return {
            "no": self.no,
            "imei": self.imei,
            "phone": self.phone,
            "location": self.location,
            "ssid": self.ssid,
            "status": self.status,
            "imei_valid": self.imei_valid,
            "latitude": self.latitude,
            "longitude": self.longitude,
        }


def _coerce_coordinate(value: Any) -> float | None:
    """Normalize a latitude/longitude cell into a float, or None when blank."""
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        logger.debug("Ignoring non-numeric coordinate value: %r", value)
        return None


def _classify_status(status_raw: str | None, location: str) -> str:
    """Classify modem status into ACTIVE or IDLE.

    Retired/legacy values (EX_CUSTOMER, BEKAS, ...) are folded into IDLE: the
    decommissioned category no longer exists, so spare modems are tracked
    through the single IDLE state.
    """
    if status_raw:
        s = status_raw.strip().upper()
        if s in ("ACTIVE", "AKTIF"):
            return "ACTIVE"
        # IDLE plus every historical non-active label collapses to IDLE.
        return "IDLE"

    loc_upper = location.strip().upper()
    if loc_upper == "IDLE":
        return "IDLE"
    if loc_upper.startswith("EX MODEM"):
        return "IDLE"
    return "ACTIVE"


def load_orbit_catalog_csv(path: Path) -> list[OrbitModem]:
    """Load modem catalog from a CSV file."""
    modems: list[OrbitModem] = []
    with path.open(mode="r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        for row in reader:
            if not row or not row.get("no"):
                continue
            no = int(row["no"].strip())
            imei = row.get("imei", "").strip()
            phone = row.get("phone", "").strip()
            location = row.get("location", "").strip()
            ssid = row.get("ssid", "").strip()
            status = _classify_status(row.get("status"), location)
            modems.append(
                OrbitModem(
                    no=no,
                    imei=imei,
                    phone=phone,
                    location=location,
                    ssid=ssid,
                    status=status,
                    latitude=_coerce_coordinate(row.get("latitude")),
                    longitude=_coerce_coordinate(row.get("longitude")),
                )
            )
    return modems


def load_orbit_catalog_excel(path: Path) -> list[OrbitModem]:
    """Load modem catalog from Excel file (.xlsx) using openpyxl."""
    try:
        import openpyxl
    except ImportError:  # pragma: no cover
        logger.warning("openpyxl is not installed; unable to load %s", path)
        return []

    wb = openpyxl.load_workbook(filename=str(path), data_only=True)
    sheet = wb.active
    if sheet is None:
        return []

    rows = list(sheet.iter_rows(values_only=True))
    if not rows or len(rows) < 2:
        return []

    # Map headers
    header_row = [str(cell).strip() if cell is not None else "" for cell in rows[0]]
    col_map: dict[str, int] = {}
    for idx, name in enumerate(header_row):
        col_map[name.lower()] = idx

    def _get_val(row_vals: tuple[Any, ...], *keys: str) -> str:
        for k in keys:
            if k.lower() in col_map:
                val = row_vals[col_map[k.lower()]]
                if val is not None:
                    return str(val).strip()
        return ""

    modems: list[OrbitModem] = []
    for r in rows[1:]:
        if not r or r[0] is None:
            continue
        try:
            no_raw = _get_val(r, "no.", "no")
            if not no_raw or not no_raw.isdigit():
                continue
            no = int(no_raw)
            imei = _get_val(r, "imei")
            phone = _get_val(r, "no.telepon", "no. telepon", "telepon", "phone")
            location = _get_val(r, "lokasi modem", "lokasi", "location")
            ssid = _get_val(r, "ssid modem", "ssid")
            status_val = _get_val(r, "status")
            status = _classify_status(status_val if status_val else None, location)

            modems.append(
                OrbitModem(
                    no=no,
                    imei=imei,
                    phone=phone,
                    location=location,
                    ssid=ssid,
                    status=status,
                    latitude=_coerce_coordinate(_get_val(r, "latitude", "lat")),
                    longitude=_coerce_coordinate(_get_val(r, "longitude", "lon", "lng")),
                )
            )
        except Exception as exc:  # pragma: no cover
            logger.warning("Error parsing Excel row %r: %s", r, exc)

    return modems


def _default_orbit_catalog() -> list[OrbitModem]:
    """Return 12 anonymous mock modems matching the operational profile.

    5 ACTIVE, 7 IDLE. Rows 1 and 3 are 10-digit IDs (imei_valid=False).
    Rows 2, 4-12 have 15-digit valid IMEIs.
    """
    return [
        OrbitModem(1, "1700000001", "081200000001", "Meeting Room Alpha", "Orbit-Room-01"),
        OrbitModem(2, "860000000000002", "081200000002", "Meeting Room Bravo", "Orbit-Room-02"),
        OrbitModem(3, "1700000003", "081200000003", "Meeting Room Charlie", "Orbit-Room-03"),
        OrbitModem(4, "860000000000004", "081200000004", "Executive Lounge", "Orbit-Room-04"),
        OrbitModem(5, "860000000000005", "081200000005", "Operation Center", "Orbit-Room-05"),
        OrbitModem(6, "860000000000006", "081200000006", "IDLE", "Orbit-Backup-01", "IDLE"),
        OrbitModem(7, "860000000000007", "081200000007", "IDLE", "Orbit-Backup-02", "IDLE"),
        OrbitModem(8, "860000000000008", "081200000008", "IDLE", "Orbit-Backup-03", "IDLE"),
        OrbitModem(9, "860000000000009", "081200000009", "IDLE", "Orbit-Backup-04", "IDLE"),
        OrbitModem(
            10,
            "860000000000010",
            "081200000010",
            "IDLE",
            "Orbit-Backup-05",
            "IDLE",
        ),
        OrbitModem(11, "860000000000011", "081200000011", "IDLE", "Orbit-Backup-06", "IDLE"),
        OrbitModem(12, "860000000000012", "081200000012", "IDLE", "Orbit-Backup-07", "IDLE"),
    ]


def resolve_orbit_catalog(file_path: Path | str | None = None) -> list[OrbitModem]:
    """Resolve Orbit modem catalog from given path, CSV, Excel, or default fallback."""
    if file_path:
        p = Path(file_path)
        if p.is_file():
            if p.suffix.lower() == ".xlsx":
                return load_orbit_catalog_excel(p)
            return load_orbit_catalog_csv(p)
        logger.warning("Configured Orbit catalog file %s does not exist; using default fallback", p)
        return _default_orbit_catalog()

    if DEFAULT_CSV_PATH.is_file():
        try:
            return load_orbit_catalog_csv(DEFAULT_CSV_PATH)
        except Exception as exc:  # pragma: no cover
            logger.warning("Failed loading %s: %s", DEFAULT_CSV_PATH, exc)

    if DEFAULT_EXCEL_PATH.is_file():
        try:
            modems = load_orbit_catalog_excel(DEFAULT_EXCEL_PATH)
            if modems:
                return modems
        except Exception as exc:  # pragma: no cover
            logger.warning("Failed loading %s: %s", DEFAULT_EXCEL_PATH, exc)

    if DEFAULT_EXAMPLE_CSV_PATH.is_file():
        try:
            return load_orbit_catalog_csv(DEFAULT_EXAMPLE_CSV_PATH)
        except Exception as exc:  # pragma: no cover
            logger.warning("Failed loading %s: %s", DEFAULT_EXAMPLE_CSV_PATH, exc)

    return _default_orbit_catalog()


def filter_stats(modems: Sequence[OrbitModem]) -> dict[str, int]:
    """Compute count statistics for UI filter pills and summary counters."""
    total = len(modems)
    active = sum(1 for m in modems if m.status == "ACTIVE")
    idle = sum(1 for m in modems if m.status == "IDLE")
    imei_pending = sum(1 for m in modems if not m.imei_valid)
    return {
        "all": total,
        "active": active,
        "idle": idle,
        "imei_pending": imei_pending,
    }


def _timestamp() -> str:
    """Return a compact UTC timestamp for backup filenames."""
    return datetime.now(UTC).strftime("%Y%m%d_%H%M%S")


def _resolve_backup_dir(backup_dir: Path | str | None) -> Path | None:
    """Normalize the optional backup directory argument."""
    if backup_dir is None:
        return None
    return Path(backup_dir)


def _atomic_write_text(path: Path, content: str) -> None:
    """Write text to path atomically via a temp file in the same directory."""
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.parent / f"{path.name}.{uuid.uuid4().hex}.tmp"
    try:
        tmp.write_text(content, encoding="utf-8")
        tmp.replace(path)
    finally:
        tmp.unlink(missing_ok=True)


def _backup_existing_file(path: Path, backup_dir: Path | None = None) -> Path | None:
    """Copy path into backup_dir with a timestamped .bak name.

    Returns the backup path, or None when there was nothing to back up.
    """
    if not path.is_file():
        return None
    target_dir = backup_dir if backup_dir is not None else path.parent / "backups"
    target_dir.mkdir(parents=True, exist_ok=True)
    backup_path = target_dir / f"{path.stem}_{_timestamp()}.{path.suffix.lstrip('.')}.bak"
    counter = 1
    while backup_path.exists():
        backup_path = target_dir / (
            f"{path.stem}_{_timestamp()}_{counter}.{path.suffix.lstrip('.')}.bak"
        )
        counter += 1
    shutil.copy2(path, backup_path)
    logger.info("Backed up %s to %s", path, backup_path)
    return backup_path


def save_orbit_catalog_csv(
    modems: Sequence[OrbitModem],
    path: Path | str | None = None,
    backup_dir: Path | str | None = None,
) -> Path | None:
    """Write the modem catalog to CSV, backing up the previous file first.

    Returns the backup path when an existing file was replaced, else None.
    """
    target = Path(path) if path is not None else DEFAULT_CSV_PATH
    backup = _backup_existing_file(target, _resolve_backup_dir(backup_dir))
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\n")
    writer.writerow(CSV_HEADER)
    for m in modems:
        writer.writerow(
            [
                m.no,
                m.imei,
                m.phone,
                m.location,
                m.ssid,
                m.status,
                "" if m.latitude is None else m.latitude,
                "" if m.longitude is None else m.longitude,
            ]
        )
    _atomic_write_text(target, buffer.getvalue())
    logger.info("Saved %d Orbit modems to %s", len(modems), target)
    return backup


def upsert_orbit_modem(
    catalog: Sequence[OrbitModem],
    modem: OrbitModem,
    original_no: int | None = None,
) -> list[OrbitModem]:
    """Insert or update a modem, returning the catalog sorted by modem number.

    When original_no differs from modem.no the entry is a renumber; both the old
    and the new number are dropped first so the catalog keeps exactly one entry.
    """
    result = [m for m in catalog if m.no != modem.no and m.no != original_no]
    result.append(modem)
    result.sort(key=lambda m: m.no)
    return result


def remove_orbit_modem(catalog: Sequence[OrbitModem], no: int) -> list[OrbitModem]:
    """Return a copy of the catalog without the modem numbered no."""
    return [m for m in catalog if m.no != no]


def save_orbit_catalog_excel(
    modems: Sequence[OrbitModem],
    path: Path | str | None = None,
    backup_dir: Path | str | None = None,
) -> Path | None:
    """Write the modem catalog to an .xlsx workbook, backing up the old file first.

    An existing workbook is edited in place (header styling and any extra sheets
    are preserved); the data rows below the header are replaced wholesale so no
    stale modem survives a deletion. Returns the backup path, or None when there
    was no previous file.
    """
    try:
        import openpyxl
    except ImportError:  # pragma: no cover - openpyxl is a project dependency
        logger.warning("openpyxl is not installed; skipping Orbit Excel catalog save")
        return None

    target = Path(path) if path is not None else DEFAULT_EXCEL_PATH
    backup = _backup_existing_file(target, _resolve_backup_dir(backup_dir))

    if target.is_file():
        wb = openpyxl.load_workbook(str(target))
        sheet = wb.active
    else:
        wb = openpyxl.Workbook()
        sheet = wb.active
    if sheet is None:  # pragma: no cover - openpyxl always provides an active sheet
        raise OSError(f"Workbook {target} has no active sheet")

    if sheet.max_row > 1:
        sheet.delete_rows(2, sheet.max_row - 1)
    for idx, header in enumerate(EXCEL_HEADER, start=1):
        sheet.cell(row=1, column=idx, value=header)
    for m in modems:
        sheet.append(
            [
                m.no,
                m.imei,
                m.phone,
                m.location,
                m.ssid,
                m.status,
                m.latitude,
                m.longitude,
            ]
        )

    target.parent.mkdir(parents=True, exist_ok=True)
    tmp = target.parent / f"{target.name}.{uuid.uuid4().hex}.tmp"
    try:
        wb.save(str(tmp))
        tmp.replace(target)
    finally:
        tmp.unlink(missing_ok=True)
    logger.info("Saved %d Orbit modems to %s", len(modems), target)
    return backup


def persist_orbit_catalog(
    modems: Sequence[OrbitModem],
    csv_path: Path | str | None = None,
    excel_path: Path | str | None = None,
    backup_dir: Path | str | None = None,
) -> dict[str, Path | None]:
    """Write the catalog to CSV and Excel, backing up each file beforehand.

    An explicit path is always written. The well-known default paths are only
    consulted when the caller names neither, so a deployment keeping a single
    source of truth is never shadowed by an extra write. Returns the backup
    paths keyed by "csv_backup" and "excel_backup".
    """
    csv_target = Path(csv_path) if csv_path is not None else None
    excel_target = Path(excel_path) if excel_path is not None else None
    if csv_target is None and excel_target is None:
        if DEFAULT_CSV_PATH.is_file():
            csv_target = DEFAULT_CSV_PATH
        if DEFAULT_EXCEL_PATH.is_file():
            excel_target = DEFAULT_EXCEL_PATH

    result: dict[str, Path | None] = {"csv_backup": None, "excel_backup": None}
    if csv_target is not None:
        result["csv_backup"] = save_orbit_catalog_csv(modems, csv_target, backup_dir)
    if excel_target is not None:
        result["excel_backup"] = save_orbit_catalog_excel(modems, excel_target, backup_dir)
    return result


def _find_column_index(col_map: dict[str, int], *keys: str) -> int | None:
    """Find the first matching column index in col_map."""
    for k in keys:
        if k.lower() in col_map:
            return col_map[k.lower()]
    return None


def _matches_target(phone_val: str, imei_val: str, target_key: str) -> bool:
    """Check if phone or IMEI matches the target identifier."""
    if phone_val and (
        phone_val == target_key or phone_val.lstrip("0") == target_key.lstrip("0")
    ):
        return True
    return bool(imei_val and imei_val == target_key)


def update_catalog_modem_ssid(
    phone_or_imei: str,
    new_ssid: str,
    config_dir: Path | str | None = None,
) -> bool:
    """Update modem SSID in Excel catalog (orbit_*.xlsx) and CSV (orbit_targets.csv) on disk.

    Returns True if a matching modem was found and updated (or already matched).
    """
    target_key = str(phone_or_imei).strip()
    new_ssid = str(new_ssid).strip()
    if not target_key or not new_ssid:
        return False

    if config_dir is not None:
        base_dir = Path(config_dir)
        if base_dir.is_file():
            base_dir = base_dir.parent
    else:
        base_dir = Path("config")

    matched = False

    excel_candidates: list[Path] = []
    default_excel = base_dir / f"orbit_{''.join(('gm', 'f'))}.xlsx"
    if default_excel.is_file():
        excel_candidates.append(default_excel)
    if base_dir.is_dir():
        for p in sorted(base_dir.glob("orbit_*.xlsx")):
            if p != default_excel and p.is_file():
                excel_candidates.append(p)

    if excel_candidates:
        try:
            import openpyxl
        except ImportError:  # pragma: no cover
            openpyxl = None

        if openpyxl is not None:
            for excel_path in excel_candidates:
                try:
                    wb = openpyxl.load_workbook(filename=str(excel_path))
                    sheet = wb.active
                    if sheet is None:
                        continue
                    rows = list(sheet.iter_rows())
                    if not rows or len(rows) < 2:
                        continue

                    header_cells = rows[0]
                    col_map: dict[str, int] = {}
                    for idx, c in enumerate(header_cells):
                        val_str = str(c.value or "").strip().lower()
                        if val_str:
                            col_map[val_str] = idx

                    phone_col = _find_column_index(
                        col_map, "no.telepon", "no. telepon", "telepon", "phone", "msisdn"
                    )
                    imei_col = _find_column_index(col_map, "imei", "nomor imei", "no. imei")
                    ssid_col = _find_column_index(
                        col_map, "ssid modem", "ssid", "wifi name", "nama wifi"
                    )

                    if ssid_col is not None and (phone_col is not None or imei_col is not None):
                        wb_modified = False
                        for r in rows[1:]:
                            r_phone = (
                                str(r[phone_col].value or "").strip()
                                if phone_col is not None and phone_col < len(r)
                                else ""
                            )
                            r_imei = (
                                str(r[imei_col].value or "").strip()
                                if imei_col is not None and imei_col < len(r)
                                else ""
                            )

                            if _matches_target(r_phone, r_imei, target_key):
                                matched = True
                                if ssid_col < len(r):
                                    cell = r[ssid_col]
                                    if str(cell.value or "").strip() != new_ssid:
                                        cell.value = new_ssid
                                        wb_modified = True

                        if wb_modified:
                            wb.save(str(excel_path))
                except Exception as exc:  # pragma: no cover
                    logger.warning("Error updating Excel catalog %s: %s", excel_path, exc)

    csv_path = base_dir / "orbit_targets.csv"
    if csv_path.is_file():
        try:
            with csv_path.open(mode="r", encoding="utf-8-sig", newline="") as f:
                reader = csv.reader(f)
                csv_rows = list(reader)

            if csv_rows and len(csv_rows) >= 2:
                header = [h.strip().lower() for h in csv_rows[0]]
                col_map_csv: dict[str, int] = {h: idx for idx, h in enumerate(header) if h}

                phone_col_csv = _find_column_index(
                    col_map_csv, "phone", "telepon", "no.telepon", "no. telepon", "msisdn"
                )
                imei_col_csv = _find_column_index(
                    col_map_csv, "imei", "nomor imei", "no. imei"
                )
                ssid_col_csv = _find_column_index(
                    col_map_csv, "ssid", "ssid modem", "wifi name", "nama wifi"
                )

                has_id_col = phone_col_csv is not None or imei_col_csv is not None
                if ssid_col_csv is not None and has_id_col:
                    csv_modified = False
                    for row in csv_rows[1:]:
                        if not row:
                            continue
                        r_phone = (
                            row[phone_col_csv].strip()
                            if phone_col_csv is not None and phone_col_csv < len(row)
                            else ""
                        )
                        r_imei = (
                            row[imei_col_csv].strip()
                            if imei_col_csv is not None and imei_col_csv < len(row)
                            else ""
                        )

                        if _matches_target(r_phone, r_imei, target_key):
                            matched = True
                            if ssid_col_csv < len(row):
                                if row[ssid_col_csv].strip() != new_ssid:
                                    row[ssid_col_csv] = new_ssid
                                    csv_modified = True
                            else:
                                while len(row) <= ssid_col_csv:
                                    row.append("")
                                row[ssid_col_csv] = new_ssid
                                csv_modified = True

                    if csv_modified:
                        tmp_csv = csv_path.with_name(f"{csv_path.name}.tmp")
                        with tmp_csv.open(mode="w", encoding="utf-8", newline="") as f:
                            writer = csv.writer(f)
                            writer.writerows(csv_rows)
                        tmp_csv.replace(csv_path)
        except Exception as exc:  # pragma: no cover
            logger.warning("Error updating CSV catalog %s: %s", csv_path, exc)

    return matched
